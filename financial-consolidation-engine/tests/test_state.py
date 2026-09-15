"""
Tests for financial consolidation workbook.
"""

import pytest
import os
import sys
from openpyxl import load_workbook


# ---- Helpers ----

def get_sheet(wb, target):
    """Get sheet by name, case/space/underscore/hyphen insensitive."""
    norm = target.lower().replace(" ", "").replace("_", "").replace("-", "")
    for sn in wb.sheetnames:
        if sn.lower().replace(" ", "").replace("_", "").replace("-", "") == norm:
            return wb[sn]
    return None


def find_sheet_strict(wb, target):
    """Find sheet by exact name, then fallback to case-insensitive."""
    if target in wb.sheetnames:
        return wb[target]
    lower = target.lower()
    for sn in wb.sheetnames:
        if sn.lower() == lower:
            return wb[sn]
    raise KeyError(f"Sheet '{target}' not found. Available: {wb.sheetnames}")


def get_all_numbers(ws):
    """Extract all numeric values from a worksheet."""
    nums = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, (int, float)) and cell is not None:
                nums.append(float(cell))
    return nums


def get_all_text(ws):
    """Extract all text values (lowercased) from a worksheet."""
    texts = []
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and cell.strip():
                texts.append(cell.strip().lower())
    return texts


def find_row_values(ws, label):
    """Find numeric values in the row containing the given label."""
    label_lower = label.lower().strip()
    for row in ws.iter_rows(values_only=True):
        has_label = False
        for cell in row:
            if isinstance(cell, str) and cell.strip().lower() == label_lower:
                has_label = True
                break
        if has_label:
            return [float(c) for c in row
                    if isinstance(c, (int, float)) and c is not None]
    return []


def approx(actual, expected, rel_tol=0.005, abs_tol=1.0):
    """Check if actual is approximately equal to expected."""
    if expected == 0:
        return abs(actual) < abs_tol
    return abs(actual - expected) <= max(abs(expected) * rel_tol, abs_tol)


def has_approx(values, expected, rel_tol=0.005, abs_tol=1.0):
    """Check if any value in list approximately matches expected."""
    for v in values:
        if approx(v, expected, rel_tol, abs_tol):
            return True
    return False


def compute_pmt(pv, rate_monthly, nper):
    """Standard PMT formula."""
    if rate_monthly == 0:
        return pv / nper
    return pv * rate_monthly / (1 - (1 + rate_monthly) ** (-nper))


def compute_remaining_balance(pv, rate_monthly, pmt, k):
    """Remaining balance after k payments."""
    if rate_monthly == 0:
        return pv - pmt * k
    return (pv * (1 + rate_monthly) ** k
            - pmt * ((1 + rate_monthly) ** k - 1) / rate_monthly)


# ---- Compute all expected values from source data ----

@pytest.fixture(scope="module")
def source_data():
    """Parse source workbook and independently compute expected values."""
    src = load_workbook("/app/group_financials.xlsx", data_only=True)

    print(f"Source workbook sheets: {src.sheetnames}", file=sys.stderr)

    # Exchange rates
    fx = {}
    ws = find_sheet_strict(src, "ExchangeRates")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1] is not None:
            fx[str(row[0])] = float(row[1])

    # Ownership stakes
    ownership = {}
    ws = find_sheet_strict(src, "Ownership")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            ownership[str(row[0])] = float(row[1])

    # Aggregate subsidiary financials in USD
    revenue_by_cat = {}
    expense_by_cat = {}
    sub_ni = {}
    for name in ["SubA", "SubB", "SubC", "SubD"]:
        ws = find_sheet_strict(src, name)
        s_rev = 0
        s_exp = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            type_, cat, amt, cur = (str(row[0]), str(row[1]),
                                    float(row[2]), str(row[3]))
            usd_amt = amt * fx[cur]
            if type_ == "Revenue":
                revenue_by_cat[cat] = revenue_by_cat.get(cat, 0) + usd_amt
                s_rev += usd_amt
            else:
                expense_by_cat[cat] = expense_by_cat.get(cat, 0) + usd_amt
                s_exp += usd_amt
        sub_ni[name] = s_rev - s_exp

    total_gross_rev = sum(revenue_by_cat.values())
    total_gross_exp = sum(expense_by_cat.values())

    # Intercompany eliminations
    ic_rev_elim = {}
    ic_exp_elim = {}
    total_ic = 0.0
    ws = find_sheet_strict(src, "Intercompany")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        amt_usd = float(row[2]) * fx[str(row[3])]
        rev_cat, exp_cat = str(row[4]), str(row[5])
        ic_rev_elim[rev_cat] = ic_rev_elim.get(rev_cat, 0) + amt_usd
        ic_exp_elim[exp_cat] = ic_exp_elim.get(exp_cat, 0) + amt_usd
        total_ic += amt_usd

    net_rev = total_gross_rev - total_ic
    net_exp = total_gross_exp - total_ic
    net_income = net_rev - net_exp

    # NCI calculation
    total_nci = sum(sub_ni[n] * (1 - ownership[n])
                    for n in sub_ni if ownership[n] < 1.0)
    parent_ni = net_income - total_nci

    # Loan amortization
    loans = []
    ws = find_sheet_strict(src, "Loans")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        pv = float(row[1])
        annual_rate = float(row[2])
        term = int(row[3])
        paid = int(row[4])
        r = annual_rate / 12
        pmt = compute_pmt(pv, r, term)
        remaining = compute_remaining_balance(pv, r, pmt, paid)
        principal_paid = pv - remaining
        interest_paid = pmt * paid - principal_paid
        remaining_interest = pmt * (term - paid) - remaining
        loans.append({
            "id": str(row[0]), "pmt": pmt, "remaining": remaining,
            "interest_paid": interest_paid, "principal_paid": principal_paid,
            "remaining_interest": remaining_interest,
        })

    # Ratios
    net_cogs = expense_by_cat.get("COGS", 0) - ic_exp_elim.get("COGS", 0)
    gross_margin = (net_rev - net_cogs) / net_rev
    operating_margin = net_income / net_rev
    ic_pct = total_ic / total_gross_rev
    expense_ratio = net_exp / net_rev
    cogs_ratio = net_cogs / net_rev

    # Break-even
    fixed_costs = sum(expense_by_cat.get(c, 0)
                      for c in ["Salaries", "R&D", "Depreciation"])
    variable_costs = net_cogs + expense_by_cat.get("Marketing", 0)
    var_cost_ratio = variable_costs / net_rev
    contribution_margin = 1 - var_cost_ratio
    breakeven_rev = fixed_costs / contribution_margin

    # Sensitivity: recompute NI at different FX factors
    def ni_at_factor(f):
        adj_fx = {cur: (rate if cur == "USD" else rate * f)
                  for cur, rate in fx.items()}
        adj_rev = 0
        adj_exp = 0
        for sname in ["SubA", "SubB", "SubC", "SubD"]:
            sub_ws = find_sheet_strict(src, sname)
            for row in sub_ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:
                    continue
                amt_usd = float(row[2]) * adj_fx[str(row[3])]
                if str(row[0]) == "Revenue":
                    adj_rev += amt_usd
                else:
                    adj_exp += amt_usd
        return adj_rev - adj_exp

    sensitivity = {
        "minus_20": ni_at_factor(0.8),
        "minus_10": ni_at_factor(0.9),
        "plus_10": ni_at_factor(1.1),
        "plus_20": ni_at_factor(1.2),
    }

    return {
        "total_gross_rev": total_gross_rev,
        "total_gross_exp": total_gross_exp,
        "total_ic": total_ic,
        "net_rev": net_rev,
        "net_exp": net_exp,
        "net_income": net_income,
        "revenue_by_cat": revenue_by_cat,
        "expense_by_cat": expense_by_cat,
        "ic_rev_elim": ic_rev_elim,
        "ic_exp_elim": ic_exp_elim,
        "net_cogs": net_cogs,
        "loans": loans,
        "gross_margin": gross_margin,
        "operating_margin": operating_margin,
        "ic_pct": ic_pct,
        "expense_ratio": expense_ratio,
        "cogs_ratio": cogs_ratio,
        "fixed_costs": fixed_costs,
        "variable_costs": variable_costs,
        "var_cost_ratio": var_cost_ratio,
        "contribution_margin": contribution_margin,
        "breakeven_rev": breakeven_rev,
        "total_nci": total_nci,
        "parent_ni": parent_ni,
        "sensitivity": sensitivity,
    }


@pytest.fixture(scope="module")
def output_wb():
    """Load the agent's output workbook."""
    path = "/app/consolidated.xlsx"
    assert os.path.exists(path), f"Output workbook not found at {path}"
    return load_workbook(path, data_only=True)


# ==== Sheet Existence ====

class TestSheetExistence:
    def test_consolidated_pl_exists(self, output_wb):
        assert get_sheet(output_wb, "ConsolidatedPL") is not None

    def test_amortization_exists(self, output_wb):
        assert get_sheet(output_wb, "Amortization") is not None

    def test_ratios_exists(self, output_wb):
        assert get_sheet(output_wb, "Ratios") is not None

    def test_breakeven_exists(self, output_wb):
        assert get_sheet(output_wb, "BreakEven") is not None

    def test_sensitivity_exists(self, output_wb):
        assert get_sheet(output_wb, "Sensitivity") is not None


# ==== ConsolidatedPL ====

class TestConsolidatedPL:
    def _nums(self, output_wb):
        return get_all_numbers(get_sheet(output_wb, "ConsolidatedPL"))

    def test_gross_revenue(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["total_gross_rev"]), \
            f"Expected gross revenue ~{source_data['total_gross_rev']:.0f}"

    def test_ic_elimination_amount(self, output_wb, source_data):
        nums = self._nums(output_wb)
        ic = source_data["total_ic"]
        assert has_approx(nums, ic) or has_approx(nums, -ic), \
            f"Expected IC elimination ~{ic:.0f}"

    def test_net_revenue(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["net_rev"]), \
            f"Expected net revenue ~{source_data['net_rev']:.0f}"

    def test_gross_expenses(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["total_gross_exp"]), \
            f"Expected gross expenses ~{source_data['total_gross_exp']:.0f}"

    def test_net_expenses(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["net_exp"]), \
            f"Expected net expenses ~{source_data['net_exp']:.0f}"

    def test_net_income(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["net_income"]), \
            f"Expected net income ~{source_data['net_income']:.0f}"

    def test_revenue_labels(self, output_wb, source_data):
        texts = get_all_text(get_sheet(output_wb, "ConsolidatedPL"))
        for cat in ["product sales", "service revenue"]:
            assert cat in texts, f"Missing label '{cat}'"

    def test_expense_labels(self, output_wb, source_data):
        texts = get_all_text(get_sheet(output_wb, "ConsolidatedPL"))
        for cat in ["cogs", "salaries"]:
            assert cat in texts, f"Missing label '{cat}'"

    def test_net_cogs(self, output_wb, source_data):
        ws = get_sheet(output_wb, "ConsolidatedPL")
        row_vals = find_row_values(ws, "COGS")
        assert has_approx(row_vals, source_data["net_cogs"]), \
            f"Expected net COGS ~{source_data['net_cogs']:.0f} in COGS row"

    def test_gross_cogs(self, output_wb, source_data):
        ws = get_sheet(output_wb, "ConsolidatedPL")
        row_vals = find_row_values(ws, "COGS")
        expected = source_data["expense_by_cat"]["COGS"]
        assert has_approx(row_vals, expected), \
            f"Expected gross COGS ~{expected:.0f} in COGS row"

    def test_product_sales_net(self, output_wb, source_data):
        ws = get_sheet(output_wb, "ConsolidatedPL")
        row_vals = find_row_values(ws, "Product Sales")
        gross = source_data["revenue_by_cat"]["Product Sales"]
        elim = source_data["ic_rev_elim"].get("Product Sales", 0)
        expected_net = gross - elim
        assert has_approx(row_vals, expected_net), \
            f"Expected net Product Sales ~{expected_net:.0f}"


# ==== NCI ====

class TestNCI:
    def _nums(self, output_wb):
        return get_all_numbers(get_sheet(output_wb, "ConsolidatedPL"))

    def test_nci_value(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["total_nci"]), \
            f"Expected NCI ~{source_data['total_nci']:.0f}"

    def test_parent_net_income(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb), source_data["parent_ni"]), \
            f"Expected parent NI ~{source_data['parent_ni']:.0f}"

    def test_nci_label(self, output_wb):
        texts = get_all_text(get_sheet(output_wb, "ConsolidatedPL"))
        found = any("non-controlling" in t or "nci" in t or "minority" in t
                     for t in texts)
        assert found, "No NCI / non-controlling interest label found"


# ==== Amortization ====

class TestAmortization:
    def _nums(self, output_wb):
        return get_all_numbers(get_sheet(output_wb, "Amortization"))

    def test_loan1_pmt(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][0]["pmt"]), \
            f"Loan 1 PMT ~{source_data['loans'][0]['pmt']:.2f} not found"

    def test_loan2_pmt(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][1]["pmt"]), \
            f"Loan 2 PMT ~{source_data['loans'][1]['pmt']:.2f} not found"

    def test_loan3_pmt(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][2]["pmt"]), \
            f"Loan 3 PMT ~{source_data['loans'][2]['pmt']:.2f} not found"

    def test_loan1_remaining_balance(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][0]["remaining"]), \
            f"Loan 1 balance ~{source_data['loans'][0]['remaining']:.2f} not found"

    def test_loan2_remaining_balance(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][1]["remaining"]), \
            f"Loan 2 balance ~{source_data['loans'][1]['remaining']:.2f} not found"

    def test_loan3_remaining_balance(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][2]["remaining"]), \
            f"Loan 3 balance ~{source_data['loans'][2]['remaining']:.2f} not found"

    def test_loan1_interest_paid(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][0]["interest_paid"]), \
            f"Loan 1 interest ~{source_data['loans'][0]['interest_paid']:.2f} not found"

    def test_loan2_interest_paid(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][1]["interest_paid"]), \
            f"Loan 2 interest ~{source_data['loans'][1]['interest_paid']:.2f} not found"

    def test_loan3_interest_paid(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["loans"][2]["interest_paid"]), \
            f"Loan 3 interest ~{source_data['loans'][2]['interest_paid']:.2f} not found"

    def test_loan_ids_present(self, output_wb, source_data):
        texts = get_all_text(get_sheet(output_wb, "Amortization"))
        for loan in source_data["loans"]:
            lid = loan["id"].lower()
            assert lid in texts, f"Loan ID '{loan['id']}' not found"


# ==== Ratios ====

class TestRatios:
    def _nums(self, output_wb):
        return get_all_numbers(get_sheet(output_wb, "Ratios"))

    def test_gross_margin(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["gross_margin"]
        assert (has_approx(nums, exp, rel_tol=0.01, abs_tol=0.005)
                or has_approx(nums, exp * 100, rel_tol=0.01, abs_tol=0.5)), \
            f"Gross margin ~{exp:.4f} not found"

    def test_operating_margin(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["operating_margin"]
        assert (has_approx(nums, exp, rel_tol=0.01, abs_tol=0.005)
                or has_approx(nums, exp * 100, rel_tol=0.01, abs_tol=0.5)), \
            f"Operating margin ~{exp:.4f} not found"

    def test_ic_revenue_pct(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["ic_pct"]
        assert (has_approx(nums, exp, rel_tol=0.01, abs_tol=0.005)
                or has_approx(nums, exp * 100, rel_tol=0.01, abs_tol=0.5)), \
            f"IC revenue pct ~{exp:.4f} not found"

    def test_cogs_ratio(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["cogs_ratio"]
        assert (has_approx(nums, exp, rel_tol=0.01, abs_tol=0.005)
                or has_approx(nums, exp * 100, rel_tol=0.01, abs_tol=0.5)), \
            f"COGS ratio ~{exp:.4f} not found"

    def test_expense_ratio(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["expense_ratio"]
        assert (has_approx(nums, exp, rel_tol=0.01, abs_tol=0.005)
                or has_approx(nums, exp * 100, rel_tol=0.01, abs_tol=0.5)), \
            f"Expense ratio ~{exp:.4f} not found"


# ==== Break-Even ====

class TestBreakEven:
    def _nums(self, output_wb):
        return get_all_numbers(get_sheet(output_wb, "BreakEven"))

    def test_fixed_costs(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["fixed_costs"]), \
            f"Fixed costs ~{source_data['fixed_costs']:.0f} not found"

    def test_variable_costs(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["variable_costs"]), \
            f"Variable costs ~{source_data['variable_costs']:.0f} not found"

    def test_contribution_margin(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["contribution_margin"]
        assert (has_approx(nums, exp, rel_tol=0.01, abs_tol=0.005)
                or has_approx(nums, exp * 100, rel_tol=0.01, abs_tol=0.5)), \
            f"Contribution margin ~{exp:.4f} not found"

    def test_breakeven_revenue(self, output_wb, source_data):
        assert has_approx(self._nums(output_wb),
                          source_data["breakeven_rev"]), \
            f"Break-even revenue ~{source_data['breakeven_rev']:.0f} not found"


# ==== Sensitivity ====

class TestSensitivity:
    def _nums(self, output_wb):
        return get_all_numbers(get_sheet(output_wb, "Sensitivity"))

    def test_minus_20_scenario(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["sensitivity"]["minus_20"]
        assert has_approx(nums, exp), \
            f"-20% scenario NI ~{exp:.0f} not found"

    def test_minus_10_scenario(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["sensitivity"]["minus_10"]
        assert has_approx(nums, exp), \
            f"-10% scenario NI ~{exp:.0f} not found"

    def test_plus_10_scenario(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["sensitivity"]["plus_10"]
        assert has_approx(nums, exp), \
            f"+10% scenario NI ~{exp:.0f} not found"

    def test_plus_20_scenario(self, output_wb, source_data):
        nums = self._nums(output_wb)
        exp = source_data["sensitivity"]["plus_20"]
        assert has_approx(nums, exp), \
            f"+20% scenario NI ~{exp:.0f} not found"
