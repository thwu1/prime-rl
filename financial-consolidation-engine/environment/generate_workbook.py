#!/usr/bin/env python3
"""Generate the source financial workbook for Nexus Global Holdings."""

import sys
from openpyxl import Workbook, load_workbook

OUTPUT_PATH = "/app/group_financials.xlsx"

EXPECTED_SHEETS = [
    "Config", "ExchangeRates",
    "SubA", "SubB", "SubC", "SubD",
    "Intercompany", "Loans", "Ownership",
]

wb = Workbook()

# --- Config ---
ws = wb.active
ws.title = "Config"
ws.append(["Parameter", "Value"])
ws.append(["Company", "Nexus Global Holdings"])
ws.append(["Reporting_Currency", "USD"])
ws.append(["Fiscal_Year", 2025])

# --- Exchange Rates ---
ws = wb.create_sheet("ExchangeRates")
ws.append(["Currency", "Rate_to_USD"])
for cur, rate in [("USD", 1.0), ("EUR", 1.085), ("JPY", 0.006667), ("BRL", 0.2)]:
    ws.append([cur, rate])

# --- Subsidiaries ---
subsidiaries = {
    "SubA": ("USD", [
        ("Revenue", "Product Sales", 5200000),
        ("Revenue", "Service Revenue", 1800000),
        ("Revenue", "Licensing", 950000),
        ("Expense", "COGS", 2600000),
        ("Expense", "Salaries", 1500000),
        ("Expense", "Marketing", 400000),
        ("Expense", "R&D", 600000),
        ("Expense", "Depreciation", 250000),
    ]),
    "SubB": ("EUR", [
        ("Revenue", "Product Sales", 3800000),
        ("Revenue", "Service Revenue", 1200000),
        ("Revenue", "Consulting", 600000),
        ("Expense", "COGS", 1900000),
        ("Expense", "Salaries", 1100000),
        ("Expense", "Marketing", 350000),
        ("Expense", "R&D", 450000),
        ("Expense", "Depreciation", 180000),
    ]),
    "SubC": ("JPY", [
        ("Revenue", "Product Sales", 620000000),
        ("Revenue", "Service Revenue", 180000000),
        ("Revenue", "Licensing", 95000000),
        ("Expense", "COGS", 340000000),
        ("Expense", "Salaries", 165000000),
        ("Expense", "Marketing", 55000000),
        ("Expense", "R&D", 80000000),
        ("Expense", "Depreciation", 28000000),
    ]),
    "SubD": ("BRL", [
        ("Revenue", "Product Sales", 12500000),
        ("Revenue", "Service Revenue", 4800000),
        ("Expense", "COGS", 7200000),
        ("Expense", "Salaries", 3600000),
        ("Expense", "Marketing", 1100000),
        ("Expense", "R&D", 800000),
        ("Expense", "Depreciation", 450000),
    ]),
}

for name in ["SubA", "SubB", "SubC", "SubD"]:
    currency, items = subsidiaries[name]
    ws = wb.create_sheet(name)
    ws.append(["Type", "Category", "Amount", "Currency"])
    for type_, cat, amt in items:
        ws.append([type_, cat, amt, currency])

# --- Intercompany Transactions ---
ws = wb.create_sheet("Intercompany")
ws.append(["From_Entity", "To_Entity", "Amount", "Currency",
           "Revenue_Category", "Expense_Category"])
ws.append(["SubA", "SubB", 200000, "USD", "Service Revenue", "COGS"])
ws.append(["SubB", "SubC", 150000, "EUR", "Consulting", "COGS"])
ws.append(["SubA", "SubD", 100000, "USD", "Licensing", "COGS"])
ws.append(["SubC", "SubA", 45000000, "JPY", "Product Sales", "COGS"])

# --- Loan Portfolio ---
ws = wb.create_sheet("Loans")
ws.append(["Loan_ID", "Principal", "Annual_Rate", "Term_Months", "Payments_Made"])
ws.append(["LOAN-001", 2000000, 0.055, 60, 24])
ws.append(["LOAN-002", 5000000, 0.042, 120, 36])
ws.append(["LOAN-003", 800000, 0.068, 36, 12])

# --- Ownership ---
ws = wb.create_sheet("Ownership")
ws.append(["Subsidiary", "Ownership_Pct"])
ws.append(["SubA", 1.0])
ws.append(["SubB", 1.0])
ws.append(["SubC", 0.85])
ws.append(["SubD", 0.70])

# --- Save ---
print(f"Sheets before save: {wb.sheetnames}")
wb.save(OUTPUT_PATH)
wb.close()

# --- Verify the saved workbook ---
verify = load_workbook(OUTPUT_PATH, data_only=True)
actual = verify.sheetnames
print(f"Sheets after reload: {actual}")
for expected in EXPECTED_SHEETS:
    if expected not in actual:
        print(f"FATAL: Missing sheet '{expected}' in saved workbook!", file=sys.stderr)
        sys.exit(1)
# Verify Ownership has data rows
ownership_ws = verify["Ownership"]
rows = list(ownership_ws.iter_rows(min_row=2, values_only=True))
if len(rows) < 4:
    print(f"FATAL: Ownership sheet has only {len(rows)} data rows, expected 4", file=sys.stderr)
    sys.exit(1)
verify.close()
print(f"Source workbook created and verified at {OUTPUT_PATH}")
