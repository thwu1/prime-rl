
import pytest
import os
from openpyxl import load_workbook

WORKBOOK_PATH = '/app/output/consolidated_report.xlsx'


def normalize(formula):
    """Normalize formula for comparison: uppercase, strip spaces."""
    if not isinstance(formula, str):
        return ""
    return formula.upper().replace(" ", "")


def assert_formula(cell, *required_keywords):
    """Assert cell value is a formula containing all required keywords (case-insensitive)."""
    val = cell.value
    assert isinstance(val, str) and val.startswith('='), \
        f"Cell {cell.coordinate}: expected formula starting with '=', got {val!r} (type={type(val).__name__})"
    norm = normalize(val)
    for kw in required_keywords:
        assert kw.upper().replace(" ", "") in norm, \
            f"Cell {cell.coordinate}: formula {val!r} missing keyword '{kw}'"


@pytest.fixture(scope='module')
def wb():
    assert os.path.exists(WORKBOOK_PATH), f"Workbook not found at {WORKBOOK_PATH}"
    return load_workbook(WORKBOOK_PATH, data_only=False)


# ---------------------------------------------------------------------------
# Structure tests
# ---------------------------------------------------------------------------
class TestWorkbookStructure:
    def test_workbook_exists(self):
        assert os.path.exists(WORKBOOK_PATH), "Output workbook does not exist"

    def test_sheet_names(self, wb):
        expected = {'Transactions', 'Products', 'Employees',
                    'RegionalSummary', 'FinancialModel', 'Dashboard'}
        actual = set(wb.sheetnames)
        missing = expected - actual
        assert not missing, f"Missing sheets: {missing}"

    def test_transaction_row_count(self, wb):
        ws = wb['Transactions']
        # 50 data rows + 1 header
        assert ws.max_row >= 51, f"Transactions: expected >=51 rows, got {ws.max_row}"

    def test_product_row_count(self, wb):
        ws = wb['Products']
        # 10 products + 1 header
        assert ws.max_row >= 11, f"Products: expected >=11 rows, got {ws.max_row}"

    def test_employee_row_count(self, wb):
        ws = wb['Employees']
        # 15 employees + 1 header
        assert ws.max_row >= 16, f"Employees: expected >=16 rows, got {ws.max_row}"


# ---------------------------------------------------------------------------
# Transaction formula tests
# ---------------------------------------------------------------------------
class TestTransactionFormulas:
    def test_revenue_is_multiplication(self, wb):
        ws = wb['Transactions']
        val = ws['G2'].value
        assert isinstance(val, str) and val.startswith('='), \
            f"G2: expected formula, got {val!r}"
        norm = normalize(val)
        assert '*' in norm, \
            f"G2: revenue formula should use multiplication, got {val!r}"
        assert 'D2' in norm and 'E2' in norm, \
            f"G2: revenue formula should reference D2 and E2, got {val!r}"

    def test_cost_formula_xlookup(self, wb):
        ws = wb['Transactions']
        assert_formula(ws['H2'], 'XLOOKUP', 'Products')

    def test_profit_formula(self, wb):
        ws = wb['Transactions']
        assert_formula(ws['I2'], 'G2', 'H2')

    def test_margin_formula(self, wb):
        ws = wb['Transactions']
        assert_formula(ws['J2'], 'I2', 'G2')

    def test_quarter_formula(self, wb):
        ws = wb['Transactions']
        assert_formula(ws['K2'], 'MONTH')

    def test_quarter_formula_has_division(self, wb):
        ws = wb['Transactions']
        val = ws['K2'].value
        assert isinstance(val, str), f"K2: expected formula, got {val!r}"
        norm = normalize(val)
        assert 'INT' in norm or '/' in norm, \
            f"K2: quarter formula should derive quarter from month, got {val!r}"

    def test_formulas_span_all_rows(self, wb):
        ws = wb['Transactions']
        last = ws.max_row
        val_g = ws[f'G{last}'].value
        assert isinstance(val_g, str) and '*' in val_g, \
            f"G{last}: expected multiplication formula, got {val_g!r}"
        assert_formula(ws[f'H{last}'], 'XLOOKUP', 'Products')


# ---------------------------------------------------------------------------
# Product formula tests
# ---------------------------------------------------------------------------
class TestProductFormulas:
    def test_total_units_sold(self, wb):
        ws = wb['Products']
        assert_formula(ws['F2'], 'SUMIFS', 'Transactions')

    def test_total_revenue(self, wb):
        ws = wb['Products']
        assert_formula(ws['G2'], 'SUMIFS', 'Transactions')

    def test_transaction_count(self, wb):
        ws = wb['Products']
        assert_formula(ws['H2'], 'COUNTIFS', 'Transactions')

    def test_avg_order_size(self, wb):
        ws = wb['Products']
        assert_formula(ws['I2'], 'IF', 'H2')

    def test_profitability_tier(self, wb):
        ws = wb['Products']
        assert_formula(ws['J2'], 'IFS')


# ---------------------------------------------------------------------------
# Employee formula tests
# ---------------------------------------------------------------------------
class TestEmployeeFormulas:
    def test_tenure_formula(self, wb):
        ws = wb['Employees']
        assert_formula(ws['G2'], 'TODAY')

    def test_salary_band(self, wb):
        ws = wb['Employees']
        assert_formula(ws['H2'], 'IFS')

    def test_bonus_formula(self, wb):
        ws = wb['Employees']
        assert_formula(ws['I2'], 'IF', 'AND')


# ---------------------------------------------------------------------------
# RegionalSummary tests
# ---------------------------------------------------------------------------
class TestRegionalSummary:
    def test_revenue_sumifs(self, wb):
        ws = wb['RegionalSummary']
        assert_formula(ws['B2'], 'SUMIFS', 'Transactions')

    def test_all_regions_present(self, wb):
        ws = wb['RegionalSummary']
        regions = set()
        for r in range(2, 7):
            v = ws.cell(row=r, column=1).value
            if v:
                regions.add(v)
        expected = {'North', 'South', 'East', 'West', 'Central'}
        assert expected == regions, f"Expected regions {expected}, got {regions}"

    def test_total_column(self, wb):
        ws = wb['RegionalSummary']
        assert_formula(ws['F2'], 'SUM')

    def test_averageifs_present(self, wb):
        ws = wb['RegionalSummary']
        found = False
        for row in range(7, 16):
            cell = ws.cell(row=row, column=2)
            if isinstance(cell.value, str) and 'AVERAGEIFS' in cell.value.upper():
                found = True
                break
        assert found, "No AVERAGEIFS formula found in RegionalSummary"

    def test_countifs_present(self, wb):
        ws = wb['RegionalSummary']
        found = False
        for row in range(7, 16):
            cell = ws.cell(row=row, column=2)
            if isinstance(cell.value, str) and 'COUNTIFS' in cell.value.upper():
                found = True
                break
        assert found, "No COUNTIFS formula found in RegionalSummary"

    def test_maxifs_present(self, wb):
        ws = wb['RegionalSummary']
        found = False
        for row in range(7, 16):
            cell = ws.cell(row=row, column=2)
            if isinstance(cell.value, str) and 'MAXIFS' in cell.value.upper():
                found = True
                break
        assert found, "No MAXIFS formula found in RegionalSummary"

    def test_minifs_present(self, wb):
        ws = wb['RegionalSummary']
        found = False
        for row in range(7, 16):
            cell = ws.cell(row=row, column=2)
            if isinstance(cell.value, str) and 'MINIFS' in cell.value.upper():
                found = True
                break
        assert found, "No MINIFS formula found in RegionalSummary"


# ---------------------------------------------------------------------------
# FinancialModel tests
# ---------------------------------------------------------------------------
class TestFinancialModel:
    def test_npv_formula(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B17'], 'NPV')

    def test_irr_formula(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B18'], 'IRR')

    def test_pmt_formula(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B19'], 'PMT')

    def test_fv_formula(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B22'], 'FV')

    def test_npv_uses_named_range(self, wb):
        ws = wb['FinancialModel']
        val = ws['B17'].value
        assert isinstance(val, str), "B17 should be a formula string"
        assert 'DiscountRate' in val, \
            f"NPV formula should reference named range 'DiscountRate', got: {val}"

    def test_initial_investment_formula(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B9'], 'B2')

    def test_total_loan_cost(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B20'], 'B19')

    def test_total_interest(self, wb):
        ws = wb['FinancialModel']
        assert_formula(ws['B21'], 'B20', 'B4')


# ---------------------------------------------------------------------------
# Named range tests
# ---------------------------------------------------------------------------
class TestNamedRanges:
    def test_discount_rate(self, wb):
        assert 'DiscountRate' in wb.defined_names, \
            "Named range 'DiscountRate' not found"

    def test_initial_investment(self, wb):
        assert 'InitialInvestment' in wb.defined_names, \
            "Named range 'InitialInvestment' not found"

    def test_cash_flows(self, wb):
        assert 'CashFlows' in wb.defined_names, \
            "Named range 'CashFlows' not found"

    def test_transaction_revenue(self, wb):
        assert 'TransactionRevenue' in wb.defined_names, \
            "Named range 'TransactionRevenue' not found"

    def test_product_catalog(self, wb):
        assert 'ProductCatalog' in wb.defined_names, \
            "Named range 'ProductCatalog' not found"


# ---------------------------------------------------------------------------
# Conditional formatting tests
# ---------------------------------------------------------------------------
class TestConditionalFormatting:
    def test_transactions_has_cf(self, wb):
        ws = wb['Transactions']
        total_rules = sum(len(cf.rules) for cf in ws.conditional_formatting)
        assert total_rules >= 2, \
            f"Transactions should have >=2 conditional formatting rules, found {total_rules}"

    def test_dashboard_has_cf(self, wb):
        ws = wb['Dashboard']
        total_rules = sum(len(cf.rules) for cf in ws.conditional_formatting)
        assert total_rules >= 1, \
            f"Dashboard should have >=1 conditional formatting rule, found {total_rules}"


# ---------------------------------------------------------------------------
# Data validation tests
# ---------------------------------------------------------------------------
class TestDataValidation:
    def test_financial_model_has_validation(self, wb):
        ws = wb['FinancialModel']
        dv_count = len(ws.data_validations.dataValidation)
        assert dv_count >= 2, \
            f"FinancialModel should have >=2 data validation rules, found {dv_count}"


# ---------------------------------------------------------------------------
# Dashboard tests
# ---------------------------------------------------------------------------
class TestDashboard:
    def test_top_product_formula(self, wb):
        ws = wb['Dashboard']
        assert_formula(ws['B3'], 'INDEX', 'MATCH', 'MAX', 'Products')

    def test_total_revenue_formula(self, wb):
        ws = wb['Dashboard']
        assert_formula(ws['B8'], 'SUM', 'Transactions')

    def test_total_transactions_formula(self, wb):
        ws = wb['Dashboard']
        assert_formula(ws['B9'], 'COUNTA', 'Transactions')

    def test_cross_sheet_regional_reference(self, wb):
        ws = wb['Dashboard']
        found = False
        for row in range(10, 16):
            cell = ws.cell(row=row, column=2)
            if isinstance(cell.value, str) and 'RegionalSummary' in cell.value:
                found = True
                break
        assert found, "Dashboard should have a cross-sheet reference to RegionalSummary"


# ---------------------------------------------------------------------------
# Anti-hardcoding tests
# ---------------------------------------------------------------------------
class TestNoHardcodedValues:
    def test_revenue_column_has_formulas(self, wb):
        ws = wb['Transactions']
        for row in range(2, 12):
            val = ws.cell(row=row, column=7).value
            assert isinstance(val, str) and val.startswith('='), \
                f"G{row}: expected formula, got {val!r} (type={type(val).__name__})"

    def test_cost_column_has_formulas(self, wb):
        ws = wb['Transactions']
        for row in range(2, 12):
            val = ws.cell(row=row, column=8).value
            assert isinstance(val, str) and val.startswith('='), \
                f"H{row}: expected formula, got {val!r}"

    def test_product_aggregation_has_formulas(self, wb):
        ws = wb['Products']
        for row in range(2, 6):
            val = ws.cell(row=row, column=6).value
            assert isinstance(val, str) and val.startswith('='), \
                f"Products F{row}: expected formula, got {val!r}"

    def test_employee_tenure_has_formulas(self, wb):
        ws = wb['Employees']
        for row in range(2, 6):
            val = ws.cell(row=row, column=7).value
            assert isinstance(val, str) and val.startswith('='), \
                f"Employees G{row}: expected formula, got {val!r}"

    def test_dashboard_top_product_is_formula(self, wb):
        ws = wb['Dashboard']
        val = ws['B3'].value
        assert isinstance(val, str) and val.startswith('='), \
            f"Dashboard B3: expected formula, got {val!r} — must not be hardcoded"

    def test_dashboard_revenue_is_formula(self, wb):
        ws = wb['Dashboard']
        val = ws['B8'].value
        assert isinstance(val, str) and val.startswith('='), \
            f"Dashboard B8: expected formula, got {val!r} — must not be hardcoded"

    def test_irr_is_formula(self, wb):
        ws = wb['FinancialModel']
        val = ws['B18'].value
        assert isinstance(val, str) and val.startswith('='), \
            f"FinancialModel B18 (IRR): expected formula, got {val!r} — must not be hardcoded"

    def test_fv_is_formula(self, wb):
        ws = wb['FinancialModel']
        val = ws['B22'].value
        assert isinstance(val, str) and val.startswith('='), \
            f"FinancialModel B22 (FV): expected formula, got {val!r} — must not be hardcoded"


# ---------------------------------------------------------------------------
# Data quality tests — verify data was cleaned
# ---------------------------------------------------------------------------
class TestDataQuality:
    def test_regions_normalized_in_transactions(self, wb):
        ws = wb['Transactions']
        valid_regions = {'North', 'South', 'East', 'West', 'Central'}
        for row in range(2, min(ws.max_row + 1, 52)):
            val = ws.cell(row=row, column=3).value
            assert val in valid_regions, \
                f"Transactions C{row}: region '{val}' not normalized to title case"

    def test_dates_are_datetime_objects(self, wb):
        ws = wb['Transactions']
        from datetime import datetime
        for row in range(2, min(ws.max_row + 1, 10)):
            val = ws.cell(row=row, column=1).value
            assert isinstance(val, datetime), \
                f"Transactions A{row}: expected datetime, got {type(val).__name__}: {val!r}"

    def test_product_costs_are_numeric(self, wb):
        ws = wb['Products']
        for row in range(2, min(ws.max_row + 1, 12)):
            val = ws.cell(row=row, column=4).value
            assert isinstance(val, (int, float)), \
                f"Products D{row}: cost should be numeric, got {type(val).__name__}: {val!r}"

    def test_employee_names_stripped(self, wb):
        ws = wb['Employees']
        for row in range(2, min(ws.max_row + 1, 17)):
            val = ws.cell(row=row, column=2).value
            if val:
                assert val == val.strip(), \
                    f"Employees B{row}: name has leading/trailing whitespace: {val!r}"

    def test_employee_salaries_are_numeric(self, wb):
        ws = wb['Employees']
        for row in range(2, min(ws.max_row + 1, 17)):
            val = ws.cell(row=row, column=6).value
            assert isinstance(val, (int, float)), \
                f"Employees F{row}: salary should be numeric, got {type(val).__name__}: {val!r}"
