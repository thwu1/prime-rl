#!/usr/bin/env python3

"""Build the corrected consolidated board report workbook."""

import csv
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


# ── Data cleaning helpers ─────────────────────────────────────────────────

def parse_date(s):
    """Parse dates in YYYY-MM-DD or M/D/YYYY formats."""
    s = s.strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    raise ValueError(f"Cannot parse date: {s!r}")


def parse_currency(s):
    """Strip currency symbols and commas, return float."""
    return float(s.strip().replace('$', '').replace(',', ''))


def normalize_region(s):
    """Normalize region name to title case."""
    return s.strip().title()


# ── Read source data ─────────────────────────────────────────────────────

def read_csv(filepath):
    with open(filepath, newline='') as f:
        reader = csv.reader(f)
        headers = next(reader)
        data = [row for row in reader]
    return headers, data


txn_h, txn_data = read_csv('/data/transactions.csv')
prod_h, prod_data = read_csv('/data/products.csv')
emp_h, emp_data = read_csv('/data/employees.csv')
bud_h, bud_data = read_csv('/data/budgets.csv')

wb = Workbook()

red_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 1: Transactions
# ═══════════════════════════════════════════════════════════════════════════

ws_txn = wb.active
ws_txn.title = "Transactions"

txn_headers = ['Date', 'ProductID', 'Region', 'Quantity', 'UnitPrice',
               'SalespersonID', 'Revenue', 'Cost', 'Profit', 'ProfitMargin', 'Quarter']
for c, h in enumerate(txn_headers, 1):
    ws_txn.cell(row=1, column=c, value=h).font = Font(bold=True)

for i, row in enumerate(txn_data, 2):
    ws_txn.cell(row=i, column=1, value=parse_date(row[0]))
    ws_txn.cell(row=i, column=1).number_format = 'YYYY-MM-DD'
    ws_txn.cell(row=i, column=2, value=row[1].strip())
    ws_txn.cell(row=i, column=3, value=normalize_region(row[2]))
    ws_txn.cell(row=i, column=4, value=int(row[3]))
    ws_txn.cell(row=i, column=5, value=parse_currency(row[4]))
    ws_txn.cell(row=i, column=6, value=row[5].strip())
    ws_txn.cell(row=i, column=7, value=f'=D{i}*E{i}')
    ws_txn.cell(row=i, column=8, value=f'=D{i}*XLOOKUP(B{i},Products!A:A,Products!D:D)')
    ws_txn.cell(row=i, column=9, value=f'=G{i}-H{i}')
    ws_txn.cell(row=i, column=10, value=f'=I{i}/G{i}')
    ws_txn.cell(row=i, column=11, value=f'=INT((MONTH(A{i})-1)/3)+1')

last_txn = len(txn_data) + 1
ws_txn.conditional_formatting.add(
    f'J2:J{last_txn}',
    FormulaRule(formula=['J2<0.1'], fill=red_fill)
)
ws_txn.conditional_formatting.add(
    f'J2:J{last_txn}',
    FormulaRule(formula=['J2>0.3'], fill=green_fill)
)

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 2: Products
# ═══════════════════════════════════════════════════════════════════════════

ws_prod = wb.create_sheet("Products")

prod_headers = ['ProductID', 'Name', 'Category', 'Cost', 'ListPrice',
                'TotalUnitsSold', 'TotalRevenue', 'TransactionCount',
                'AvgOrderSize', 'ProfitabilityTier']
for c, h in enumerate(prod_headers, 1):
    ws_prod.cell(row=1, column=c, value=h).font = Font(bold=True)

for i, row in enumerate(prod_data, 2):
    ws_prod.cell(row=i, column=1, value=row[0].strip())
    ws_prod.cell(row=i, column=2, value=row[1].strip())
    ws_prod.cell(row=i, column=3, value=row[2].strip())
    ws_prod.cell(row=i, column=4, value=parse_currency(row[3]))
    ws_prod.cell(row=i, column=5, value=parse_currency(row[4]))
    ws_prod.cell(row=i, column=6, value=f'=SUMIFS(Transactions!D:D,Transactions!B:B,A{i})')
    ws_prod.cell(row=i, column=7, value=f'=SUMIFS(Transactions!G:G,Transactions!B:B,A{i})')
    ws_prod.cell(row=i, column=8, value=f'=COUNTIFS(Transactions!B:B,A{i})')
    ws_prod.cell(row=i, column=9, value=f'=IF(H{i}>0,G{i}/H{i},0)')
    ws_prod.cell(row=i, column=10,
                 value=f'=IFS((E{i}-D{i})/E{i}>0.4,"High",(E{i}-D{i})/E{i}>0.2,"Medium",TRUE,"Low")')

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 3: Employees
# ═══════════════════════════════════════════════════════════════════════════

ws_emp = wb.create_sheet("Employees")

emp_headers = ['EmployeeID', 'Name', 'Department', 'Region', 'HireDate',
               'Salary', 'TenureYears', 'SalaryBand', 'AnnualBonus']
for c, h in enumerate(emp_headers, 1):
    ws_emp.cell(row=1, column=c, value=h).font = Font(bold=True)

for i, row in enumerate(emp_data, 2):
    ws_emp.cell(row=i, column=1, value=row[0].strip())
    ws_emp.cell(row=i, column=2, value=row[1].strip())
    ws_emp.cell(row=i, column=3, value=row[2].strip())
    ws_emp.cell(row=i, column=4, value=row[3].strip())
    ws_emp.cell(row=i, column=5, value=parse_date(row[4]))
    ws_emp.cell(row=i, column=5).number_format = 'YYYY-MM-DD'
    ws_emp.cell(row=i, column=6, value=parse_currency(row[5]))
    ws_emp.cell(row=i, column=7, value=f'=(TODAY()-E{i})/365.25')
    ws_emp.cell(row=i, column=8,
                value=f'=IFS(F{i}>=100000,"Executive",F{i}>=75000,"Senior",F{i}>=50000,"Mid",TRUE,"Junior")')
    ws_emp.cell(row=i, column=9,
                value=f'=IF(AND(G{i}>2,F{i}<80000),F{i}*0.1,IF(G{i}>5,F{i}*0.15,F{i}*0.05))')

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 4: RegionalSummary
# ═══════════════════════════════════════════════════════════════════════════

ws_reg = wb.create_sheet("RegionalSummary")

ws_reg.cell(row=1, column=1, value="Region").font = Font(bold=True)
for q in range(1, 5):
    ws_reg.cell(row=1, column=q + 1, value=q).font = Font(bold=True)
ws_reg.cell(row=1, column=6, value="Total").font = Font(bold=True)

regions = ["North", "South", "East", "West", "Central"]
for i, region in enumerate(regions, 2):
    ws_reg.cell(row=i, column=1, value=region)
    for q in range(1, 5):
        col = q + 1
        cl = get_column_letter(col)
        ws_reg.cell(row=i, column=col,
                    value=f'=SUMIFS(Transactions!G:G,Transactions!C:C,$A{i},Transactions!K:K,{cl}$1)')
    ws_reg.cell(row=i, column=6, value=f'=SUM(B{i}:E{i})')

# Metrics rows
ws_reg.cell(row=8, column=1, value="Metrics").font = Font(bold=True)
metric_labels = ["AvgRevenue", "TransactionCount", "MaxTransaction", "MinTransaction"]
metric_funcs = ["AVERAGEIFS", "COUNTIFS", "MAXIFS", "MINIFS"]
for mi, (label, func) in enumerate(zip(metric_labels, metric_funcs)):
    row = 9 + mi
    ws_reg.cell(row=row, column=1, value=label)
    for q in range(1, 5):
        col = q + 1
        cl = get_column_letter(col)
        if func == "COUNTIFS":
            ws_reg.cell(row=row, column=col,
                        value=f'=COUNTIFS(Transactions!K:K,{cl}$1)')
        else:
            ws_reg.cell(row=row, column=col,
                        value=f'={func}(Transactions!G:G,Transactions!K:K,{cl}$1)')

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 5: FinancialModel
# ═══════════════════════════════════════════════════════════════════════════

ws_fin = wb.create_sheet("FinancialModel")

ws_fin.cell(row=1, column=1, value="Parameter").font = Font(bold=True)
ws_fin.cell(row=1, column=2, value="Value").font = Font(bold=True)

params = [
    ("Initial Investment", 500000),
    ("Discount Rate", 0.08),
    ("Loan Amount", 300000),
    ("Annual Interest Rate", 0.06),
    ("Loan Term Years", 10),
]
for idx, (name, val) in enumerate(params, 2):
    ws_fin.cell(row=idx, column=1, value=name)
    ws_fin.cell(row=idx, column=2, value=val)

ws_fin.cell(row=8, column=1, value="Year").font = Font(bold=True)
ws_fin.cell(row=8, column=2, value="Cash Flow").font = Font(bold=True)

ws_fin.cell(row=9, column=1, value=0)
ws_fin.cell(row=9, column=2, value='=-B2')

cash_flows = [120000, 135000, 150000, 140000, 160000]
for idx, cf in enumerate(cash_flows):
    ws_fin.cell(row=10 + idx, column=1, value=idx + 1)
    ws_fin.cell(row=10 + idx, column=2, value=cf)

ws_fin.cell(row=16, column=1, value="Metric").font = Font(bold=True)
ws_fin.cell(row=16, column=2, value="Result").font = Font(bold=True)

metrics = [
    ("NPV", '=NPV(DiscountRate,B10:B14)+B9'),
    ("IRR", '=IRR(B9:B14)'),
    ("Monthly Payment", '=PMT(B5/12,B6*12,-B4)'),
    ("Total Loan Cost", '=B19*B6*12'),
    ("Total Interest", '=B20-B4'),
    ("Future Value 5yr", '=FV(B3,5,-AVERAGE(B10:B14))'),
]
for idx, (label, formula) in enumerate(metrics):
    ws_fin.cell(row=17 + idx, column=1, value=label)
    ws_fin.cell(row=17 + idx, column=2, value=formula)

# Data validation
dv_rate = DataValidation(type="decimal", operator="between", formula1=0, formula2=1)
dv_rate.error = "Discount rate must be between 0 and 1"
dv_rate.errorTitle = "Invalid Rate"
ws_fin.add_data_validation(dv_rate)
dv_rate.add('B3')

dv_invest = DataValidation(type="whole", operator="greaterThan", formula1=0)
dv_invest.error = "Investment must be a positive integer"
dv_invest.errorTitle = "Invalid Investment"
ws_fin.add_data_validation(dv_invest)
dv_invest.add('B2')

# ═══════════════════════════════════════════════════════════════════════════
# Sheet 6: Dashboard
# ═══════════════════════════════════════════════════════════════════════════

ws_dash = wb.create_sheet("Dashboard")

ws_dash.cell(row=1, column=1, value="Dashboard Summary").font = Font(bold=True, size=14)

dash_items = [
    (3, "Top Product by Revenue",
     '=INDEX(Products!B:B,MATCH(MAX(Products!G:G),Products!G:G,0))'),
    (4, "Top Product Revenue",
     '=MAX(Products!G:G)'),
    (5, "Lowest Product by Revenue",
     '=INDEX(Products!B:B,MATCH(MIN(Products!G:G),Products!G:G,0))'),
    (6, "Lowest Product Revenue",
     '=MIN(Products!G:G)'),
    (8, "Total Company Revenue",
     '=SUM(Transactions!G:G)'),
    (9, "Total Transactions",
     '=COUNTA(Transactions!A:A)-1'),
    (10, "Average Transaction Value",
     '=B8/B9'),
    (12, "Best Region Q1",
     '=INDEX(RegionalSummary!A:A,MATCH(MAX(RegionalSummary!B2:B6),RegionalSummary!B2:B6,0)+1)'),
    (13, "Best Region Q1 Revenue",
     '=MAX(RegionalSummary!B2:B6)'),
]
for row, label, formula in dash_items:
    ws_dash.cell(row=row, column=1, value=label)
    ws_dash.cell(row=row, column=2, value=formula)

ws_dash.conditional_formatting.add(
    'B8',
    CellIsRule(operator='greaterThan', formula=['500000'], fill=green_fill)
)
ws_dash.conditional_formatting.add(
    'B8',
    CellIsRule(operator='lessThan', formula=['200000'], fill=red_fill)
)

# ═══════════════════════════════════════════════════════════════════════════
# Named Ranges
# ═══════════════════════════════════════════════════════════════════════════

named_ranges = [
    ('DiscountRate', "FinancialModel!$B$3"),
    ('InitialInvestment', "FinancialModel!$B$2"),
    ('CashFlows', "FinancialModel!$B$10:$B$14"),
    ('TransactionRevenue', "Transactions!$G:$G"),
    ('ProductCatalog', f"Products!$A$1:$E${len(prod_data) + 1}"),
]
for name, ref in named_ranges:
    dn = DefinedName(name, attr_text=ref)
    wb.defined_names.add(dn)

# ═══════════════════════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════════════════════

os.makedirs('/app/output', exist_ok=True)
wb.save('/app/output/consolidated_report.xlsx')
print("Workbook generated: /app/output/consolidated_report.xlsx")
